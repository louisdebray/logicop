#!/usr/bin/env python3
"""Convertit Base_Produits_Manutention.xlsx en catalogue-data.json pour l'outil
"Catalogue produits". À relancer à chaque mise à jour du fichier Excel source.

Usage : python3 build-catalogue-data.py /chemin/vers/Base_Produits_Manutention.xlsx
        (par défaut, cherche le fichier dans ~/Downloads)
"""
import openpyxl
import json
import sys
import os

BRAND_SHEETS = ['Mitsubishi', 'Hangcha', 'BYD', 'EP', 'Noblelift']
COL_MAP = {
    1: 'marque', 2: 'modele', 3: 'type', 4: 'sousType', 5: 'energie',
    6: 'typeBatterie', 7: 'nbRoues', 8: 'typeMat', 9: 'refMat',
    10: 'h1', 11: 'h3', 12: 'hauteurPoseMax', 13: 'llibreSansDos',
    14: 'llibreAvecDos', 15: 'matDeployeSansDos', 16: 'matDeployeAvecDos',
    17: 'capaciteNominale', 18: 'capaciteRes500', 19: 'capaciteRes600',
    20: 'rayonBraquage', 21: 'ficheSource', 22: 'dateMaj', 23: 'remarques',
    24: 'h6', 25: 'hauteurPassageMini'
}


def clean_val(val):
    if val is None:
        return None
    if isinstance(val, str):
        s = val.strip()
        return 'NC' if s == 'NC' else s
    return val


def build(xlsx_path, out_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    models = {}
    model_order = []

    for sheet_name in BRAND_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        if ws.max_row <= 1:
            continue

        current_modele = None
        for r in range(2, ws.max_row + 1):
            row = {}
            for c, key in COL_MAP.items():
                row[key] = clean_val(ws.cell(r, c).value)

            if row['modele']:
                current_modele = row['modele']
            else:
                row['modele'] = current_modele
            if not row['marque']:
                row['marque'] = sheet_name
            if not row['modele']:
                continue
            if row['dateMaj'] is not None and not isinstance(row['dateMaj'], str):
                row['dateMaj'] = str(row['dateMaj'])

            key = row['marque'] + '|' + row['modele']
            if key not in models:
                models[key] = {
                    'marque': row['marque'], 'modele': row['modele'], 'type': row['type'],
                    'sousTypeSet': set(), 'energieSet': set(), 'typeBatterieSet': set(),
                    'nbRouesSet': set(), 'ficheSource': row['ficheSource'],
                    'dateMaj': row['dateMaj'], 'configs': []
                }
                model_order.append(key)

            md = models[key]
            if row['sousType']:
                md['sousTypeSet'].add(row['sousType'])
            if row['energie']:
                md['energieSet'].add(row['energie'])
            if row['typeBatterie']:
                md['typeBatterieSet'].add(row['typeBatterie'])
            if row['nbRoues']:
                md['nbRouesSet'].add(row['nbRoues'])

            md['configs'].append({
                'typeMat': row['typeMat'], 'refMat': row['refMat'], 'sousType': row['sousType'],
                'h1': row['h1'], 'h3': row['h3'], 'hauteurPoseMax': row['hauteurPoseMax'],
                'llibreSansDos': row['llibreSansDos'], 'llibreAvecDos': row['llibreAvecDos'],
                'matDeployeSansDos': row['matDeployeSansDos'], 'matDeployeAvecDos': row['matDeployeAvecDos'],
                'capaciteNominale': row['capaciteNominale'], 'capaciteRes500': row['capaciteRes500'],
                'capaciteRes600': row['capaciteRes600'], 'rayonBraquage': row['rayonBraquage'],
                'hauteurPassageMini': row['hauteurPassageMini'], 'h6': row['h6'],
                'remarques': row['remarques']
            })

    final_list = []
    for key in model_order:
        md = models[key]
        configs = md['configs']
        h3_vals = [c['h3'] for c in configs if isinstance(c['h3'], (int, float))]
        cap_vals = [c['capaciteNominale'] for c in configs if isinstance(c['capaciteNominale'], (int, float))]
        wa_vals = [c['rayonBraquage'] for c in configs if isinstance(c['rayonBraquage'], (int, float))]
        h6_vals = [c['h6'] for c in configs if isinstance(c['h6'], (int, float))]

        final_list.append({
            'marque': md['marque'], 'modele': md['modele'], 'type': md['type'],
            'sousType': list(md['sousTypeSet'])[0] if len(md['sousTypeSet']) == 1 else None,
            'energie': list(md['energieSet'])[0] if md['energieSet'] else None,
            'typeBatterie': list(md['typeBatterieSet'])[0] if len(md['typeBatterieSet']) == 1 else None,
            'nbRoues': list(md['nbRouesSet'])[0] if md['nbRouesSet'] else None,
            'ficheSource': md['ficheSource'], 'dateMaj': md['dateMaj'],
            'h6': h6_vals[0] if h6_vals else None,
            'h3Min': min(h3_vals) if h3_vals else None,
            'h3Max': max(h3_vals) if h3_vals else None,
            'capMin': min(cap_vals) if cap_vals else None,
            'capMax': max(cap_vals) if cap_vals else None,
            'waMin': min(wa_vals) if wa_vals else None,
            'nbConfigs': len(configs),
            'configs': configs
        })

    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(final_list, f, ensure_ascii=False, separators=(',', ':'))

    print(f"{len(final_list)} modèles écrits dans {out_path}")
    wb.close()


if __name__ == '__main__':
    xlsx = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser('~/Downloads/Base_Produits_Manutention.xlsx')
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'catalogue-data.json')
    build(xlsx, out)
